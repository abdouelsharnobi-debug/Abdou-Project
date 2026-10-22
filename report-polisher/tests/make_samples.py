"""Builds deliberately messy sample reports (docx / md / pdf) used by the tests and for a quick demo."""
import sys
from pathlib import Path

import docx
from docx.shared import Pt
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas


def make_docx(path):
    d = docx.Document()
    d.add_paragraph("pump station inspection").runs[0].bold = True
    for line in ("Author: M. Ali", "Date: 2026-05-02", "Ref: INSP-77"):
        d.add_paragraph(line)
    p = d.add_paragraph("1. INTRODUCTION")
    d.add_paragraph("this report summarises the inspection of pump P-101 carried out on site. we didn't find any leaks")
    p = d.add_paragraph(); r = p.add_run("Findings"); r.bold = True
    d.add_paragraph("Bearing temperature reached 85°C against a limit of 90°C and flow was 120m3/h .")
    d.add_paragraph("See Table 2 for details")
    d.add_paragraph("Table 2: Pump readings")
    t = d.add_table(rows=3, cols=3)
    for i, row in enumerate((("Item", "Value", "Limit"), ("Temp", "85", "90"), ("Flow", "120", "150"))):
        for j, v in enumerate(row):
            t.cell(i, j).text = v
    d.add_paragraph("pump is fine - recommend continuing operation").style = "List Bullet"
    d.add_paragraph("Recommendations")
    d.add_paragraph("Replace the bearing at next shutdown").style = "List Bullet"
    d.save(path)


def make_md(path):
    Path(path).write_text("""# gearbox failure analysis

**Author:** K. Novak
**Date:** 2026-01-15

## Summary
Gearbox GB-4 failed after 4,200 h.

## Findings
The pinion shows pitting. Oil sample showed 35 ppm iron.

```
iron = 35 ppm
limit = 20 ppm
```

## Conclusion
Lubrication was inadequate.

## Recommendations
1. Change oil every 2,000 h
2. Install an oil-condition sensor
""")


def make_pdf(path):
    c = canvas.Canvas(str(path), pagesize=A4)
    y = 800
    for line in ("Boiler Efficiency Test", "", "INTRODUCTION", "The test measured boiler B-2 efficiency at full load.", "",
                 "RESULTS", "Efficiency was 91.2 % at 100 % load and 88.4 % at 50 % load.", "", "CONCLUSION",
                 "The boiler meets the 88 % guarantee."):
        c.drawString(60, y, line)
        y -= 18
    c.save()


if __name__ == "__main__":
    out = Path(sys.argv[1])
    out.mkdir(parents=True, exist_ok=True)
    make_docx(out / "pump_inspection.docx")
    make_md(out / "gearbox.md")
    make_pdf(out / "boiler.pdf")
