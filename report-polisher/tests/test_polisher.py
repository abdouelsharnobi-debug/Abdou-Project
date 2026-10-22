import tempfile
import unittest
from pathlib import Path

import docx

from reportpolish.llm import Ollama
from reportpolish.model import Branding
from reportpolish.pipeline import convert
from reportpolish.structure import classify
from reportpolish.textfix import clean_heading, clean_inline
from tests import make_samples

TMP = Path(tempfile.mkdtemp())
make_samples.make_docx(TMP / "a.docx")
make_samples.make_md(TMP / "b.md")
make_samples.make_pdf(TMP / "c.pdf")


def headings(path):
    return [p.text for p in docx.Document(str(path)).paragraphs if p.style.name.startswith("Heading")]


class TextFix(unittest.TestCase):
    def test_units_and_contractions(self):
        self.assertEqual(clean_inline("flow was 120m3/h and 25°C, we didn't stop"), "Flow was 120 m³/h and 25 °C, we did not stop")

    def test_ids_untouched(self):
        self.assertIn("P2A", clean_inline("check valve P2A and part 3D"))

    def test_heading(self):
        self.assertEqual(clean_heading("2.1 TEST SETUP AND HVAC:", {"HVAC"}), "Test Setup and HVAC")

    def test_classify(self):
        self.assertEqual(classify("Conclusions and Recommendations"), "conclusions")
        self.assertEqual(classify("Test Results"), "findings")
        self.assertIsNone(classify("Plant description"))


class Pipeline(unittest.TestCase):
    def run_one(self, name):
        return convert(TMP / name, TMP / "out", Branding(company="Acme"), want_pdf=False)

    def test_docx_structure_and_metadata(self):
        r = self.run_one("a.docx")
        m = r["report"].meta
        self.assertEqual((m["title"], m["author"], m["docno"]), ("Pump Station Inspection", "M. Ali", "INSP-77"))
        h = headings(r["docx"])
        self.assertEqual([x.split("  ")[-1] for x in h], ["Introduction", "Findings", "Recommendations"])
        self.assertTrue(r["notes"].read_text().count("Conclusions"))

    def test_markdown_code_and_summary(self):
        r = self.run_one("b.md")
        h = headings(r["docx"])
        self.assertEqual(h[0], "Executive Summary")
        self.assertEqual(h[-2:], ["2  Conclusions", "3  Recommendations"])

    def test_pdf_input(self):
        r = self.run_one("c.pdf")
        self.assertEqual(r["report"].meta["title"], "Boiler Efficiency Test")

    def test_sections_reordered(self):
        src = TMP / "d.md"
        src.write_text("# T\n\n## Recommendations\nDo it now please.\n\n## Introduction\nWhy we wrote this report today.\n\n## Findings\nIt broke.\n")
        h = headings(convert(src, TMP / "out", want_pdf=False)["docx"])
        self.assertEqual([x.split("  ")[-1] for x in h], ["Introduction", "Findings", "Recommendations"])

    def test_table_references_renumbered(self):
        src = TMP / "e.md"
        src.write_text("# T\n\n## Results\nSee Table 7 for data.\n\nTable 7: Data\n| a | b |\n|---|---|\n| 1 | 2 |\n")
        r = convert(src, TMP / "out", want_pdf=False)
        self.assertIn("See Table 1 for data.", [p.text for p in docx.Document(str(r["docx"])).paragraphs])


class AI(unittest.TestCase):
    def test_numbers_must_survive(self):
        para = "The pump ran for 40 hours at 1500 rpm without any trouble at all."
        self.assertEqual(Ollama("x", generate=lambda p: "The pump ran fine for 41 hours at 1500 rpm.").polish(para), para)
        good = "The pump operated for 40 hours at 1500 rpm without any trouble."
        self.assertEqual(Ollama("x", generate=lambda p: good).polish(para), good)

    def test_offline_model_is_harmless(self):
        def boom(p):
            raise OSError("no server")
        para = "The pump ran for 40 hours at 1500 rpm without any trouble at all."
        self.assertEqual(Ollama("x", generate=boom).polish(para), para)


if __name__ == "__main__":
    unittest.main()
