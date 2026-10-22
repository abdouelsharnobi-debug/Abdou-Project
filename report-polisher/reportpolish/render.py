"""Render a structured Report to a professionally formatted Word document (+ PDF through LibreOffice)."""
import io
import re
import shutil
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .model import Branding

FONT = "Calibri"
MONO = "Consolas"
TEXT_W = Cm(16.4)
_INLINE = re.compile(r"(\*\*\*.+?\*\*\*|\*\*.+?\*\*|(?<![\w*])\*(?=\S)[^*]+?(?<=\S)\*(?![\w*])|`[^`]+`)")


def _rgb(hexs):
    return RGBColor.from_string(hexs.upper())


def _font(run, name=FONT, size=None, bold=None, italic=None, color=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), name)
    if rf.getparent() is None:
        rpr.insert(0, rf)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = _rgb(color)


def _shade(cell, hexs):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexs)
    tcpr.append(shd)


def _borders(table, color="BFBFBF", sz=4, inside=True):
    tblpr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right") + (("insideH", "insideV") if inside else ()):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), str(sz))
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), color)
        b.append(e)
    tblpr.append(b)


def _cell_margins(table, top=50, bottom=50, left=90, right=90):
    m = OxmlElement("w:tblCellMar")
    for k, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement(f"w:{k}")
        e.set(qn("w:w"), str(v))
        e.set(qn("w:type"), "dxa")
        m.append(e)
    table._tbl.tblPr.append(m)


def _para_border(par, edge="bottom", color="1F3864", sz=8, space=1):
    ppr = par._p.get_or_add_pPr()
    pb = ppr.find(qn("w:pBdr"))
    if pb is None:
        pb = OxmlElement("w:pBdr")
        ppr.append(pb)
    e = OxmlElement(f"w:{edge}")
    e.set(qn("w:val"), "single")
    e.set(qn("w:sz"), str(sz))
    e.set(qn("w:space"), str(space))
    e.set(qn("w:color"), color)
    pb.append(e)


def _field(par, instr, text="1", **fmt):
    for kind in ("begin", None, "separate", "text", "end"):
        r = par.add_run()
        _font(r, **fmt)
        if kind == "text":
            r.text = text
        elif kind is None:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = f" {instr} "
            r._r.append(it)
        else:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), kind)
            r._r.append(fc)


def add_inline(par, text, size=None, color=None, base_bold=None):
    for part in _INLINE.split(text):
        if not part:
            continue
        r = par.add_run()
        if part.startswith("***") and part.endswith("***") and len(part) > 6:
            r.text, kw = part[3:-3], dict(bold=True, italic=True)
        elif part.startswith("**") and part.endswith("**") and len(part) > 4:
            r.text, kw = part[2:-2], dict(bold=True)
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            r.text = part[1:-1]
            _font(r, MONO, (size or 10.5) - 1, color=color)
            continue
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            r.text, kw = part[1:-1], dict(italic=True)
        else:
            r.text, kw = part, {}
        if base_bold and "bold" not in kw:
            kw["bold"] = True
        _font(r, size=size, color=color, **kw)


def _styles(doc, accent):
    st = doc.styles
    n = st["Normal"]
    n.font.name, n.font.size = FONT, Pt(10.5)
    n.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    n.paragraph_format.space_after = Pt(6)
    n.paragraph_format.line_spacing = 1.15
    for name, size, before, after in (("Heading 1", 16, 20, 8), ("Heading 2", 13, 14, 5), ("Heading 3", 11.5, 10, 4)):
        s = st[name]
        s.font.name, s.font.size, s.font.bold, s.font.italic = FONT, Pt(size), True, False
        s.font.color.rgb = _rgb(accent)
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            s.element.rPr.rFonts.set(qn(a), FONT)
        for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            if s.element.rPr.rFonts.get(qn(a)) is not None:
                del s.element.rPr.rFonts.attrib[qn(a)]
        s.paragraph_format.space_before, s.paragraph_format.space_after = Pt(before), Pt(after)
        s.paragraph_format.keep_with_next = True
    c = st["Caption"]
    c.font.name, c.font.size, c.font.bold, c.font.italic = FONT, Pt(9), True, False
    c.font.color.rgb = _rgb("404040")
    c.paragraph_format.space_after, c.paragraph_format.space_before = Pt(4), Pt(8)
    lb = st["List Bullet"]
    lb.paragraph_format.space_after = Pt(3)
    for i, ind in enumerate((0, 0.6, 1.2), 1):
        t = st.add_style(f"toc {i}", 1)
        t.base_style = n
        t.font.size, t.font.bold = Pt(11 if i == 1 else 10), i == 1
        t.paragraph_format.left_indent = Cm(ind)
        t.paragraph_format.space_after = Pt(3 if i > 1 else 4)
        t.paragraph_format.space_before = Pt(6 if i == 1 else 0)
        t.paragraph_format.tab_stops.add_tab_stop(TEXT_W, WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)


def _cover(doc, meta, brand):
    a = brand.accent
    if brand.logo and Path(brand.logo).is_file():
        p = doc.add_paragraph()
        p.add_run().add_picture(brand.logo, height=Cm(1.8))
    elif brand.company:
        p = doc.add_paragraph()
        r = p.add_run(brand.company.upper())
        _font(r, size=13, bold=True, color=a)
    for _ in range(5):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    _font(p.add_run(meta["title"]), size=30, bold=True, color=a)
    _para_border(p, "bottom", a, 12, 8)
    if meta.get("subtitle") or meta.get("project"):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        _font(p.add_run(meta.get("subtitle") or meta.get("project")), size=15, color="595959")
    for _ in range(7):
        doc.add_paragraph()
    rows = [(lbl, meta[k]) for k, lbl in (("project", "Project"), ("client", "Client"), ("docno", "Document No."),
                                          ("revision", "Revision"), ("date", "Date"), ("author", "Prepared by"),
                                          ("reviewer", "Reviewed by"), ("approver", "Approved by"),
                                          ("department", "Department")) if meta.get(k)
            and not (k == "project" and meta.get("subtitle") is None and meta.get("project") == meta.get("subtitle"))]
    if rows:
        t = doc.add_table(rows=len(rows), cols=2)
        t.alignment = WD_TABLE_ALIGNMENT.LEFT
        t.autofit = False
        t.columns[0].width, t.columns[1].width = Cm(4.2), Cm(8.8)
        _borders(t, "D9D9D9", 4, inside=False)
        for i, (k, v) in enumerate(rows):
            for j, txt in enumerate((k, v)):
                c = t.cell(i, j)
                c.width = Cm(4.2) if j == 0 else Cm(8.8)
                c.paragraphs[0].paragraph_format.space_after = Pt(2)
                _font(c.paragraphs[0].add_run(txt), size=10, bold=(j == 0), color="595959" if j == 0 else None)
                if j == 0:
                    _shade(c, "F2F2F2")
    if brand.confidential:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(24)
        _font(p.add_run(brand.confidential), size=9, italic=True, color="7F7F7F")
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def _header_footer(doc, meta, brand):
    sec = doc.sections[0]
    sec.different_first_page_header_footer = True
    hp = sec.header.paragraphs[0]
    hp.style = doc.styles["Normal"]
    hp.paragraph_format.tab_stops.add_tab_stop(TEXT_W, WD_TAB_ALIGNMENT.RIGHT)
    _font(hp.add_run(meta["title"]), size=8.5, color="595959")
    right = " | ".join(x for x in (meta.get("docno"), f"Rev {meta['revision']}" if meta.get("revision") else "") if x)
    _font(hp.add_run("\t" + right), size=8.5, color="595959")
    _para_border(hp, "bottom", "BFBFBF", 4, 4)
    fp = sec.footer.paragraphs[0]
    fp.style = doc.styles["Normal"]
    fp.paragraph_format.tab_stops.add_tab_stop(TEXT_W, WD_TAB_ALIGNMENT.RIGHT)
    _para_border(fp, "top", "BFBFBF", 4, 4)
    _font(fp.add_run(brand.confidential or brand.company or ""), size=8.5, color="7F7F7F")
    _font(fp.add_run("\tPage "), size=8.5, color="7F7F7F")
    _field(fp, "PAGE", size=8.5, color="7F7F7F")
    _font(fp.add_run(" of "), size=8.5, color="7F7F7F")
    _field(fp, "NUMPAGES", size=8.5, color="7F7F7F")


def _toc(doc, entries, pages, accent):
    p = doc.add_paragraph()
    _font(p.add_run("Contents"), size=18, bold=True, color=accent)
    _para_border(p, "bottom", accent, 8, 4)
    p.paragraph_format.space_after = Pt(10)
    for i, (lvl, label) in enumerate(entries):
        par = doc.add_paragraph(style=f"toc {min(lvl, 3)}")
        if i == 0:
            for kind in ("begin", "instr", "separate"):
                r = par.add_run()
                if kind == "instr":
                    it = OxmlElement("w:instrText")
                    it.set(qn("xml:space"), "preserve")
                    it.text = ' TOC \\o "1-3" \\h \\z \\u '
                    r._r.append(it)
                else:
                    fc = OxmlElement("w:fldChar")
                    fc.set(qn("w:fldCharType"), kind)
                    r._r.append(fc)
        par.add_run(f"{label}\t{pages.get(label, '')}")
        if i == len(entries) - 1:
            r = par.add_run()
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), "end")
            r._r.append(fc)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def _is_num(s):
    return bool(re.fullmatch(r"[<>≈~±+\-–]?\s*\d[\d\s.,]*\s*(%|[A-Za-z°/²³]{0,6})?", s.strip()))


def _table(doc, b, accent):
    rows = b.rows
    ncols = max(len(r) for r in rows)
    t = doc.add_table(rows=len(rows), cols=ncols)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    _borders(t)
    _cell_margins(t)
    numeric_col = [sum(_is_num(r[c]) for r in rows[1:] if r[c]) >= max(1, len([r for r in rows[1:] if r[c]]) * 0.7)
                   for c in range(ncols)]
    for i, row in enumerate(rows):
        for j in range(ncols):
            cell = t.cell(i, j)
            par = cell.paragraphs[0]
            par.paragraph_format.space_after = Pt(0)
            par.paragraph_format.line_spacing = 1.0
            txt = row[j] if j < len(row) else ""
            if i == 0:
                _font(par.add_run(txt), size=9.5, bold=True, color="FFFFFF")
                _shade(cell, accent)
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER if numeric_col[j] else WD_ALIGN_PARAGRAPH.LEFT
            else:
                add_inline(par, txt, size=9.5)
                if i % 2 == 0:
                    _shade(cell, "F2F5FA")
                if numeric_col[j]:
                    par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if i == 0:
            trpr = t.rows[0]._tr.get_or_add_trPr()
            trpr.append(OxmlElement("w:tblHeader"))
        cant = OxmlElement("w:cantSplit")
        t.rows[i]._tr.get_or_add_trPr().append(cant)
    return t


def _caption(doc, kind, b, below):
    p = doc.add_paragraph(style="Caption")
    p.paragraph_format.keep_with_next = not below
    if below:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    label = f"{kind} {b.label}"
    _font(p.add_run(label), size=9, bold=True, color="404040")
    if b.caption:
        _font(p.add_run(" – " + b.caption.rstrip(".")), size=9, bold=False, color="404040")


def build_docx(report, brand=None, pages=None):
    brand = brand or Branding()
    pages = pages or {}
    meta = report.meta
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.3)
    sec.top_margin, sec.bottom_margin = Cm(2.4), Cm(2.2)
    _styles(doc, brand.accent)
    cp = doc.core_properties
    cp.title, cp.author = meta["title"], meta.get("author", brand.company or "")
    cp.subject, cp.comments = meta.get("project", ""), "Formatted automatically by Report Polisher"

    # build heading labels first (needed for the contents list)
    entries, labels, sec_no, app_no, sub = [], {}, 0, 0, [0, 0]
    for b in report.blocks:
        if b.kind != "heading":
            continue
        key = getattr(b, "auto_key", None)
        if b.level == 1:
            if key == "summary":
                label = b.text
            elif key == "appendix" or re.match(r"appendix", b.text, re.I):
                app_no += 1
                label = b.text if re.match(r"appendix\s+\w", b.text, re.I) else f"Appendix {chr(64 + app_no)} – {b.text}"
            else:
                sec_no += 1
                sub = [0, 0]
                label = f"{sec_no}  {b.text}"
            b._numbered = key != "summary" and not (key == "appendix" or re.match(r"appendix", b.text, re.I))
        elif b.level in (2, 3) and sec_no and not app_no:
            if b.level == 2:
                sub = [sub[0] + 1, 0]
                label = f"{sec_no}.{sub[0]}  {b.text}"
            else:
                sub[1] += 1
                label = f"{sec_no}.{sub[0]}.{sub[1]}  {b.text}"
        else:
            label = b.text
        b._label = label
        if b.level <= 3:
            entries.append((b.level, label))

    _cover(doc, meta, brand)
    _header_footer(doc, meta, brand)
    if len(entries) >= 3:
        _toc(doc, entries, pages, brand.accent)

    numbered_n = 0
    for b in report.blocks:
        if b.kind == "heading":
            lvl = min(b.level, 3)
            if b.level >= 4:
                p = doc.add_paragraph()
                p.paragraph_format.keep_with_next = True
                add_inline(p, b.text, size=10.5, base_bold=True)
            else:
                doc.add_heading(b._label, level=lvl)
            numbered_n = 0
        elif b.kind == "para":
            p = doc.add_paragraph()
            add_inline(p, b.text)
            numbered_n = 0
        elif b.kind == "bullet":
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, b.text)
            numbered_n = 0
        elif b.kind == "number":
            numbered_n += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.9)
            p.paragraph_format.first_line_indent = Cm(-0.6)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.tab_stops.add_tab_stop(Cm(0.9))
            _font(p.add_run(f"{numbered_n}.\t"))
            add_inline(p, b.text)
        elif b.kind == "code":
            for line in b.text.split("\n"):
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.left_indent = Cm(0.4)
                _font(p.add_run(line or " "), MONO, 9)
                pPr = p._p.get_or_add_pPr()
                shd = OxmlElement("w:shd")
                shd.set(qn("w:val"), "clear")
                shd.set(qn("w:fill"), "F2F2F2")
                pPr.append(shd)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
        elif b.kind == "table":
            _caption(doc, "Table", b, below=False)
            _table(doc, b, brand.accent)
            doc.add_paragraph().paragraph_format.space_after = Pt(4)
        elif b.kind == "image":
            try:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.keep_with_next = True
                p.add_run().add_picture(io.BytesIO(b.data), width=Cm(14))
                _caption(doc, "Figure", b, below=True)
            except Exception:
                pass
    return doc


def soffice():
    return shutil.which("soffice") or shutil.which("libreoffice")


def to_pdf(docx_path, outdir):
    exe = soffice()
    if not exe:
        return None
    outdir = Path(outdir)
    with tempfile.TemporaryDirectory() as profile:
        subprocess.run([exe, f"-env:UserInstallation=file://{profile}", "--headless", "--convert-to", "pdf",
                        "--outdir", str(outdir), str(docx_path)], capture_output=True, timeout=180)
    pdf = outdir / (Path(docx_path).stem + ".pdf")
    return pdf if pdf.exists() else None


def _page_map(pdf, entries_labels):
    """Find the page each heading lands on so the contents list can show real page numbers."""
    try:
        from pypdf import PdfReader
        pages = [[re.sub(r"\s+", " ", l).strip() for l in (p.extract_text() or "").split("\n")] for p in PdfReader(str(pdf)).pages]
    except Exception:
        return {}
    out, cur = {}, 0
    for label in entries_labels:
        want = re.sub(r"\s+", " ", label).strip()
        for pi in range(cur, len(pages)):
            ls = pages[pi]
            if any(l == want or (i + 1 < len(ls) and f"{l} {ls[i + 1]}" == want) for i, l in enumerate(ls)):
                out[label], cur = str(pi + 1), pi
                break
    return out


def render(report, out_docx, brand=None, want_pdf=True):
    """Write DOCX (and PDF when LibreOffice is installed). Returns (docx_path, pdf_path|None)."""
    out_docx = Path(out_docx)
    pages = {}
    if soffice():
        with tempfile.TemporaryDirectory() as tmp:
            first = Path(tmp) / "pass1.docx"
            build_docx(report, brand, {}).save(first)
            pdf = to_pdf(first, tmp)
            if pdf:
                labels = [b._label for b in report.blocks if b.kind == "heading" and b.level <= 3]
                pages = _page_map(pdf, labels)
    build_docx(report, brand, pages).save(out_docx)
    pdf = to_pdf(out_docx, out_docx.parent) if want_pdf else None
    return out_docx, pdf
